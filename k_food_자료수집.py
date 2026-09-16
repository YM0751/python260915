from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen
from xml.etree import ElementTree

from docx import Document
from docx.shared import Pt


RSS_URL = (
    "https://news.google.com/rss/search?q="
    f"{quote('K-food Korean food')}&hl=ko&gl=KR&ceid=KR:ko"
)
OUTPUT_PATH = Path(__file__).with_name("k-food.docx")


def collect_news(max_articles=10):
    request = Request(RSS_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(request, timeout=15) as response:
        root = ElementTree.fromstring(response.read())

    articles = []
    for item in root.findall("./channel/item")[:max_articles]:
        articles.append(
            {
                "title": item.findtext("title", "제목 없음"),
                "link": item.findtext("link", ""),
                "published": item.findtext("pubDate", "발행일 없음"),
                "description": item.findtext("description", "내용 없음"),
            }
        )
    return articles


def create_word_file(articles):
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Malgun Gothic"
    normal_style.font.size = Pt(10)

    document.add_heading("K-food 최신 자료", level=0)
    document.add_paragraph(
        f"수집 시각: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    document.add_paragraph(
        "Google News RSS에서 'K-food Korean food'를 검색해 수집한 기사 목록입니다."
    )

    if not articles:
        document.add_paragraph("수집된 기사가 없습니다.")
    else:
        for number, article in enumerate(articles, start=1):
            document.add_heading(f"{number}. {article['title']}", level=2)
            document.add_paragraph(f"발행일: {article['published']}")
            document.add_paragraph(f"출처 링크: {article['link']}")
            document.add_paragraph(article["description"])

    document.add_heading("참고", level=1)
    document.add_paragraph(
        "기사 내용과 링크는 원문 제공처의 저작권 및 이용 약관을 준수해 활용하세요."
    )
    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    news_articles = collect_news()
    create_word_file(news_articles)
    print(f"{OUTPUT_PATH}에 기사 {len(news_articles)}건을 저장했습니다.")